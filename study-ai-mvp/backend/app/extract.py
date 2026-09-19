from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation

SUPPORTED = {".pdf", ".txt", ".md", ".docx", ".pptx"}


def extract_document(path: Path) -> tuple[str, int | None]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages: list[str] = []
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append(f"\n--- PAGE {i} ---\n{text}")
        return "\n".join(pages).strip(), len(reader.pages)

    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore"), None

    if suffix == ".docx":
        doc = DocxDocument(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip()), None

    if suffix == ".pptx":
        prs = Presentation(str(path))
        slides: list[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            bits: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    bits.append(shape.text.strip())
            slides.append(f"\n--- PAGE {i} ---\n" + "\n".join(bits))
        return "\n".join(slides).strip(), len(prs.slides)

    return "", None
